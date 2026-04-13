import openpyxl

wb = openpyxl.load_range = openpyxl.load_workbook(r"raw\2. Order form V7.xlsx", data_only=True)
sheet = wb.active

print(f"Sheet Name: {sheet.title}")
print(f"Max Row: {sheet.max_row}, Max Col: {sheet.max_column}")

for row_idx in range(1, min(20, sheet.max_row + 1)):
    row_data = []
    for col_idx in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=row_idx, column=col_idx)
        val = cell.value
        row_data.append(str(val) if val is not None else "")
    
    # only print row if it has data
    if any(row_data):
        print(f"Row {row_idx}: {row_data}")
